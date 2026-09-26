"""
Script to generate official LPJ Document in DOCX and PDF format.
Follows all academic / government formal styling guidelines:
- Paper: A4 (21.0 x 29.7 cm), Margins: Left 3.5cm, Right 3.0cm, Top 3.0cm, Bottom 3.0cm
- Font: Times New Roman, 12pt body, 1.15 line spacing, Justified, 1.0cm first line indent
- Headings: Bab 14pt Bold UPPERCASE Centered; Subbab 12pt Bold Left; Sub-subbab 12pt Bold Left
- Cover: SMP Negeri 1 Biau formal logo and typography
- Lembar Pengesahan: Formal placeholders
- TOC, TOF, TOT
- Full BAB I s.d. BAB V content from master
- Lampiran A s.d. G with tables, 25 screenshots, and 2 diagrams
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

# Ensure output directory exists
OUTPUT_DIR = os.path.abspath('docs/LPJ/final')
os.makedirs(OUTPUT_DIR, exist_ok=True)
DOCX_PATH = os.path.join(OUTPUT_DIR, 'LPJ_Pengembangan_Aplikasi_FINAL.docx')
PDF_PATH = os.path.join(OUTPUT_DIR, 'LPJ_Pengembangan_Aplikasi_FINAL.pdf')

print(f"Target DOCX: {DOCX_PATH}")
print(f"Target PDF:  {PDF_PATH}")

doc = docx.Document()

# Base style setup
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(12)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(4)
normal_style.paragraph_format.space_before = Pt(0)
normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# XML helper functions
def set_cell_border(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    cell._tc.get_or_add_tcPr().append(tcMar)

def add_field(p, text):
    run = p.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {text} </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def set_table_header_and_split(table):
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def apply_page_setup(section, is_prelim=True):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)

# Section 1: Preliminaries (i, ii, iii...)
sec1 = doc.sections[0]
apply_page_setup(sec1, is_prelim=True)
sec1.different_first_page_header_footer = True

# Prelim Page numbering format (roman lowercase)
sectPr1 = sec1._sectPr
pgNumType1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
sectPr1.append(pgNumType1)

# Header for Section 1 (shown on pages ii onwards)
hdr1 = sec1.header
hdr_p1 = hdr1.paragraphs[0]
hdr_p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hdr_p1.paragraph_format.space_after = Pt(2)
r_hdr1 = hdr_p1.add_run("APLIKASI PRESENSI SIASEK — SMP NEGERI 1 BIAU")
r_hdr1.font.name = 'Times New Roman'
r_hdr1.font.size = Pt(8.5)
r_hdr1.font.italic = True
r_hdr1.font.color.rgb = RGBColor(100, 116, 139)

# Footer for Section 1
ftr1 = sec1.footer
ftr_p1 = ftr1.paragraphs[0]
ftr_p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
r_ftr_lbl1 = ftr_p1.add_run("LAPORAN PERTANGGUNGJAWABAN\t\t")
r_ftr_lbl1.font.name = 'Times New Roman'
r_ftr_lbl1.font.size = Pt(9)
r_ftr_lbl1.font.color.rgb = RGBColor(100, 116, 139)
add_field(ftr_p1, "PAGE")
for run in ftr_p1.runs:
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 116, 139)

print("Section 1 setup completed.")
